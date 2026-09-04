"""Local document extraction and FTS retrieval for ChatMPD."""

from __future__ import annotations

import hashlib
import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

from .platform_db import PlatformDatabase


_TEXT_SUFFIXES = {
    ".txt", ".md", ".rst", ".csv", ".tsv", ".json", ".jsonl", ".toml",
    ".yaml", ".yml", ".py", ".js", ".ts", ".tsx", ".jsx", ".html", ".css",
    ".ps1", ".bat", ".cmd", ".ini", ".cfg", ".log", ".xml",
}


def _now() -> str:
    return datetime.now(UTC).isoformat()


def _fts_query(text: str) -> str:
    words = re.findall(r"[\w-]+", str(text), flags=re.UNICODE)
    return " AND ".join(f'"{word}"' for word in words)


def extract_text(path: Path, *, max_bytes: int = 32 * 1024 * 1024) -> str:
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    if source.stat().st_size > max_bytes:
        raise ValueError(f"Knowledge source is too large: {source.name}")
    suffix = source.suffix.casefold()
    if suffix in _TEXT_SUFFIXES:
        return source.read_text(encoding="utf-8", errors="replace")
    if suffix == ".docx":
        with zipfile.ZipFile(source) as archive:
            xml = archive.read("word/document.xml")
        root = ET.fromstring(xml)
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        paragraphs: list[str] = []
        for paragraph in root.findall(".//w:p", namespace):
            text = "".join(node.text or "" for node in paragraph.findall(".//w:t", namespace))
            if text.strip():
                paragraphs.append(text)
        return "\n".join(paragraphs)
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as error:
            raise RuntimeError("PDF extraction requires the local pypdf package.") from error
        reader = PdfReader(str(source))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise RuntimeError("Spreadsheet extraction requires the local openpyxl package.") from error
        workbook = load_workbook(source, read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    lines.append("\t".join(values))
        workbook.close()
        return "\n".join(lines)
    raise ValueError(f"Unsupported knowledge source type: {source.suffix or '<none>'}")


@dataclass(frozen=True)
class KnowledgeSource:
    source_id: str
    path: Path
    title: str
    kind: str
    sha256: str
    created_at: str
    updated_at: str


@dataclass(frozen=True)
class KnowledgeHit:
    source_id: str
    chunk_id: str
    path: Path
    title: str
    text: str
    score: float


class KnowledgeLibrary:
    def __init__(self, database: PlatformDatabase | None = None) -> None:
        self.database = database or PlatformDatabase()

    def ingest(self, path: Path, *, chunk_chars: int = 3000, overlap: int = 300) -> KnowledgeSource:
        source = Path(path).expanduser().resolve()
        text = extract_text(source)
        if not text.strip():
            raise ValueError(f"Knowledge source contains no readable text: {source.name}")
        if chunk_chars < 500 or overlap < 0 or overlap >= chunk_chars:
            raise ValueError("Invalid knowledge chunk settings.")
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        now = _now()
        chunks = self._chunks(text, chunk_chars=chunk_chars, overlap=overlap)
        with self.database.connect() as connection:
            existing = connection.execute(
                "SELECT source_id FROM knowledge_sources WHERE path=?", (str(source),)
            ).fetchone()
            if existing is not None:
                old_id = str(existing["source_id"])
                connection.execute("DELETE FROM knowledge_fts WHERE source_id=?", (old_id,))
                connection.execute("DELETE FROM knowledge_sources WHERE source_id=?", (old_id,))
            source_id = uuid4().hex
            connection.execute(
                "INSERT INTO knowledge_sources VALUES (?, ?, ?, ?, ?, ?, ?)",
                (source_id, str(source), source.name, source.suffix.casefold().lstrip("."), digest, now, now),
            )
            for index, chunk in enumerate(chunks):
                chunk_id = uuid4().hex
                connection.execute(
                    "INSERT INTO knowledge_chunks VALUES (?, ?, ?, ?)",
                    (chunk_id, source_id, index, chunk),
                )
                connection.execute(
                    "INSERT INTO knowledge_fts(chunk_id, source_id, text) VALUES (?, ?, ?)",
                    (chunk_id, source_id, chunk),
                )
            connection.commit()
        return self.get(source_id)

    def get(self, source_id: str) -> KnowledgeSource:
        with self.database.connect() as connection:
            row = connection.execute(
                "SELECT * FROM knowledge_sources WHERE source_id=?", (source_id,)
            ).fetchone()
        if row is None:
            raise KeyError(f"Unknown knowledge source: {source_id}")
        return self._source_from_row(row)

    def list(self) -> tuple[KnowledgeSource, ...]:
        with self.database.connect() as connection:
            rows = connection.execute(
                "SELECT * FROM knowledge_sources ORDER BY updated_at DESC"
            ).fetchall()
        return tuple(self._source_from_row(row) for row in rows)

    def search(self, text: str, *, limit: int = 12) -> tuple[KnowledgeHit, ...]:
        expression = _fts_query(text)
        if not expression:
            return ()
        with self.database.connect() as connection:
            rows = connection.execute(
                """
                SELECT f.chunk_id, f.source_id, f.text, bm25(knowledge_fts) AS rank,
                       s.path, s.title
                FROM knowledge_fts f
                JOIN knowledge_sources s ON s.source_id=f.source_id
                WHERE knowledge_fts MATCH ?
                ORDER BY rank LIMIT ?
                """,
                (expression, max(1, min(int(limit), 100))),
            ).fetchall()
        return tuple(
            KnowledgeHit(
                source_id=str(row["source_id"]),
                chunk_id=str(row["chunk_id"]),
                path=Path(str(row["path"])),
                title=str(row["title"]),
                text=str(row["text"]),
                score=float(row["rank"]),
            )
            for row in rows
        )

    @staticmethod
    def _chunks(text: str, *, chunk_chars: int, overlap: int) -> tuple[str, ...]:
        cleaned = str(text).strip()
        if not cleaned:
            return ()
        chunks: list[str] = []
        start = 0
        while start < len(cleaned):
            end = min(len(cleaned), start + chunk_chars)
            chunk = cleaned[start:end].strip()
            if chunk:
                chunks.append(chunk)
            if end >= len(cleaned):
                break
            start = max(start + 1, end - overlap)
        return tuple(chunks)

    @staticmethod
    def _source_from_row(row) -> KnowledgeSource:
        return KnowledgeSource(
            source_id=str(row["source_id"]),
            path=Path(str(row["path"])),
            title=str(row["title"]),
            kind=str(row["kind"]),
            sha256=str(row["sha256"]),
            created_at=str(row["created_at"]),
            updated_at=str(row["updated_at"]),
        )
